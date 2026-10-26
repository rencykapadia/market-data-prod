"""
The Stash (New Mexico) scraper.

Uses The Stash's own JSON exports served alongside the site. Their menu
frontend fetches these three static files on load; we do the same:

    GET https://thestashnm.com/inventory.json      -- in-stock inventory
    GET https://thestashnm.com/catalogexport.json  -- product master (strain, THC, weight)
    GET https://thestashnm.com/deals.xlsx          -- active promotions

Row shape per SKU is produced by joining inventory (source of truth for
stock and Regular Price) with the catalog (source of Strain Type, THC
Potency, Net Weight, Manufacturer) by SKU.

Locations (single Cova tenant, three stores):
    111 Lujan Rd              -> Bernalillo
    12611 Montgomery Blvd. NE -> Albuquerque
    510 HWY 528               -> Rio Rancho

Env vars:
    STASH_PROXY / HTTPS_PROXY  -- residential US proxy (required if Cloudflare 403s)

Requirements:
    pip install curl_cffi pytz pandas openpyxl requests
"""

import os
import io
import sys
import json
import time
import logging
from datetime import datetime
import pytz
import pandas as pd
from openpyxl import load_workbook

# curl_cffi impersonates Chrome's TLS fingerprint -- required to defeat
# The Stash's Cloudflare bot check on the static JSON endpoints.
try:
    from curl_cffi import requests as _http
    _HTTP_BACKEND = "curl_cffi"
except ImportError:
    import requests as _http
    _HTTP_BACKEND = "requests"

# =========================
# LOGGER
# =========================
logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
logger = logging.getLogger(__name__)

# =========================
# CONFIG
# =========================
company = "The Stash"
state = "New Mexico"

us_tz = pytz.timezone("America/Denver")   # New Mexico is Mountain Time
today = datetime.now(us_tz).strftime("%Y-%m-%d")

output_folder = "app/output/pending"
csv_folder = os.path.join(output_folder, "csv")
json_folder = os.path.join(output_folder, "json")
os.makedirs(csv_folder, exist_ok=True)
os.makedirs(json_folder, exist_ok=True)
output_file = os.path.join(csv_folder, f"{today}_{company}.csv")
json_file_path = os.path.join(json_folder, f"{today}_{company}.json")

# The Stash operates 3 NM stores; inventory.json returns all three, keyed
# by street address in the 'Location' column.
STORES = {
    "111 Lujan Rd":              "Bernalillo",
    "12611 Montgomery Blvd. NE": "Albuquerque",
    "510 HWY 528":               "Rio Rancho",
}

# Cannabis product types recognised by New Mexico's Retail Marijuana Product Type field
CANNABIS_TYPES = {"Flower", "Concentrate", "Edible"}

# Inventory Classifications that are cannabis even when NM Product Type is blank
CANNABIS_CLASSIFICATIONS = {
    "Flower", "Pre-roll", "Pre-roll Packs", "Infused Pre-roll",
    "Concentrates", "Concentrates ", "Disposable Vapes", "Vape Cartridge",
    "Edible", "Beverages", "Tincture", "Topical", "Salve", "Bath Bomb",
}

# =========================
# HTTP
# =========================
BASE_URL = "https://thestashnm.com"
HTTP_TIMEOUT = 45
REQUEST_DELAY_S = 0.2

_CHROME_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/148.0.0.0 Safari/537.36"
)

HEADERS = {
    "User-Agent": _CHROME_UA,
    "Accept": "application/json,*/*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": f"{BASE_URL}/menu.html",
    "sec-fetch-site": "same-origin",
    "sec-fetch-mode": "cors",
    "sec-fetch-dest": "empty",
}


def _resolve_proxy():
    url = (
        os.environ.get("STASH_PROXY")
        or os.environ.get("HTTPS_PROXY")
        or os.environ.get("HTTP_PROXY")
    )
    if not url:
        return None
    return {"http": url, "https": url}


def make_session():
    proxies = _resolve_proxy()
    if _HTTP_BACKEND == "curl_cffi":
        s = _http.Session(impersonate="chrome120")
    else:
        s = _http.Session()
    if proxies:
        s.proxies.update(proxies)
    logger.info(f"backend={_HTTP_BACKEND} {'via proxy' if proxies else 'DIRECT (no proxy)'}")
    return s


def _warm_up(session):
    """Hit the menu.html once to acquire Cloudflare cookies before JSON GETs."""
    try:
        session.get(f"{BASE_URL}/menu.html", headers=HEADERS, timeout=HTTP_TIMEOUT)
    except Exception as e:
        logger.warning(f"warm-up failed (continuing): {e}")


def _get(session, path):
    """GET a static asset with one 403 retry after re-warming."""
    url = f"{BASE_URL}{path}?v={int(datetime.now().timestamp()*1000)}"
    for attempt in range(2):
        r = session.get(url, headers=HEADERS, timeout=HTTP_TIMEOUT)
        if r.status_code == 403 and attempt == 0:
            logger.warning(f"403 on {path} -- re-warming and retrying")
            _warm_up(session)
            time.sleep(1)
            continue
        r.raise_for_status()
        return r
    raise RuntimeError(f"failed to GET {path}")


# =========================
# FETCHERS
# =========================
def fetch_inventory(session):
    return _get(session, "/inventory.json").json()


def fetch_catalog(session):
    return _get(session, "/catalogexport.json").json()


def fetch_deals(session):
    """
    Return {(sku, store_address): [deal_text, ...]}.
    Deals with Location='All Locations' are expanded to every configured store.
    Deals xlsx layout: row index 3 is the header, data starts at row 4.
    """
    r = _get(session, "/deals.xlsx")
    wb = load_workbook(io.BytesIO(r.content), data_only=True, read_only=True)
    ws = wb["Menu Promotion Reference"]

    all_rows = list(ws.iter_rows(values_only=True))
    if len(all_rows) < 5:
        return {}

    header = [str(c) if c is not None else "" for c in all_rows[3]]
    try:
        col_loc    = header.index("Location")
        col_sku    = header.index("SKU")
        col_promo  = header.index("Promo Price / Deal")
        col_status = header.index("Cova Status")
    except ValueError:
        logger.warning("deals.xlsx header mismatch, skipping deals")
        return {}

    deals = {}
    for row in all_rows[4:]:
        if row is None or len(row) <= col_sku:
            continue
        sku = row[col_sku]
        if not sku:
            continue
        status = str(row[col_status] or "").strip().lower()
        if status and status != "active":
            continue
        promo = row[col_promo]
        if promo is None or str(promo).strip() == "":
            continue
        loc = str(row[col_loc] or "").strip()
        text = str(promo).strip()
        sku_s = str(sku).strip()

        targets = list(STORES) if loc == "All Locations" else ([loc] if loc in STORES else [])
        for t in targets:
            bucket = deals.setdefault((sku_s, t), [])
            if text not in bucket:
                bucket.append(text)
    return deals


# =========================
# PARSER
# =========================
def is_cannabis(inv_row, cat_row):
    nm_type = (cat_row or {}).get("Retail Marijuana Product Type - New Mexico", "")
    if nm_type and nm_type in CANNABIS_TYPES:
        return True
    return (inv_row.get("Classification") or "").strip() in CANNABIS_CLASSIFICATIONS


def resolve_brand(inv_row, cat_row):
    for key in ("Manufacturer Name", "Manufacturer"):
        val = (cat_row or {}).get(key) or inv_row.get(key)
        if val:
            return str(val).strip()
    for key in ("Supplier0", "Supplier"):
        val = (cat_row or {}).get(key) or inv_row.get(key)
        if val:
            return str(val).strip()
    return ""


def resolve_weight(cat_row):
    if not cat_row:
        return ""
    nw = str(cat_row.get("Net Weight") or "").strip()
    if not nw or nw == "0":
        return ""
    unit = (cat_row.get("Unit of Measure") or "").strip()
    if unit and unit.lower() != "each":
        return f"{nw} {unit}"
    return f"{nw}g"


def resolve_thc(cat_row):
    if not cat_row:
        return ""
    for k in ("THC Potency", "THC Content"):
        v = cat_row.get(k)
        if v not in (None, ""):
            try:
                return f"{float(v):g} %"
            except (TypeError, ValueError):
                return str(v).strip()
    return ""


def format_price(v):
    if v in (None, ""):
        return ""
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return v


# =========================
# MAIN
# =========================
all_rows = []
per_store_counts = {}

try:
    session = make_session()
    logger.info(f"Starting {company} scrape | stores={len(STORES)}")

    _warm_up(session)
    logger.info("Fetching inventory.json ...")
    inventory = fetch_inventory(session)
    logger.info(f"  inventory rows: {len(inventory)}")

    logger.info("Fetching catalogexport.json ...")
    catalog = fetch_catalog(session)
    logger.info(f"  catalog rows: {len(catalog)}")

    logger.info("Fetching deals.xlsx ...")
    deals_map = fetch_deals(session)
    logger.info(f"  deals mapped: {len(deals_map)} (sku, store) pairs")

    cat_by_sku = {str(c.get("SKU", "")).strip(): c for c in catalog if c.get("SKU")}

    for store_addr, dispensary_location in STORES.items():
        store_rows_before = len(all_rows)
        logger.info(f"--- {dispensary_location} ({store_addr}) ---")

        for inv in inventory:
            try:
                qty = float(inv.get("In Stock Qty") or 0)
            except (TypeError, ValueError):
                qty = 0.0
            if qty <= 0:
                continue

            loc = str(inv.get("Location", "")).strip()
            if loc != store_addr:
                continue

            sku = str(inv.get("SKU", "")).strip()
            cat_row = cat_by_sku.get(sku)

            if not is_cannabis(inv, cat_row):
                continue

            classification = (inv.get("Classification") or "").strip()
            subcategory = ((cat_row or {}).get("Breadcrumb") or classification).strip()

            csv_row = {
                "date": today,
                "location": state,
                "company": company,
                "dispensary_location": dispensary_location,
                "name": (inv.get("Product") or "").strip(),
                "category": classification,
                "subcategory": subcategory,
                "brand": resolve_brand(inv, cat_row),
                "strain_type": ((cat_row or {}).get("Strain Type") or "").strip(),
                "THC": resolve_thc(cat_row),
                "CBD": "",
                "weight": resolve_weight(cat_row),
                "original_price": format_price(inv.get("Regular Price")),
            }
            for i, deal_text in enumerate(deals_map.get((sku, store_addr), []), 1):
                csv_row[f"deal_{i}"] = deal_text
            all_rows.append(csv_row)

        n = len(all_rows) - store_rows_before
        per_store_counts[dispensary_location] = n
        logger.info(f"=> {dispensary_location}: {n} rows")
        time.sleep(REQUEST_DELAY_S)

except Exception as e:
    # Graceful skip: a 403 with no proxy configured = this host's IP is Cloudflare-blocked. Exit 0
    # (don't fail the pipeline) and warn; set STASH_PROXY/HTTPS_PROXY to a US residential proxy to
    # enable this scraper. A 403 WITH a proxy is a real problem, so still fail then.
    if "403" in str(e) and not _resolve_proxy():
        logger.warning(
            f"⚠ {company}: SKIPPED — Cloudflare is IP-blocking this host (403) and no proxy is set. "
            f"Set STASH_PROXY or HTTPS_PROXY to a US residential proxy to enable it. No data this run."
        )
        sys.exit(0)
    logger.critical(f"❌ {company} fetch failed: {e}", exc_info=True)
    sys.exit(1)


# =========================
# SAVE JSON
# =========================
try:
    with open(json_file_path, "w", encoding="utf-8") as f:
        json.dump(
            {"per_store_counts": per_store_counts, "total_rows": len(all_rows)},
            f, indent=2,
        )
    logger.info(f"JSON saved: {json_file_path}")
except Exception as e:
    logger.critical(f"❌ JSON save failed: {e}", exc_info=True)
    sys.exit(1)


# =========================
# SAVE CSV
# =========================
try:
    df = pd.DataFrame(all_rows).drop_duplicates(
        subset=["dispensary_location", "name", "weight", "original_price"]
    )
    if df.empty or len(df) < 5:
        logger.critical(
            f"❌ Invalid CSV output. Rows: {len(df)}. "
            f"Likely 403 from Cloudflare. Set HTTPS_PROXY to a US residential proxy."
        )
        sys.exit(1)

    fixed_cols = [
        "date", "location", "company", "dispensary_location", "name",
        "category", "subcategory", "brand", "strain_type",
        "THC", "CBD", "weight", "original_price",
    ]
    deal_cols = sorted(
        [c for c in df.columns if c.startswith("deal_")],
        key=lambda x: int(x.split("_")[1]),
    )
    df = df[fixed_cols + deal_cols]
    df.to_csv(output_file, index=False, na_rep="")

    per_store = df.groupby(["dispensary_location"]).size().to_dict()
    cat_counts = df.groupby("category").size().to_dict()
    logger.info(f"CSV saved: {output_file} | Rows: {len(df)} | Deal cols: {len(deal_cols)}")
    logger.info(f"Rows per dispensary: {per_store}")
    logger.info(f"Rows per category: {cat_counts}")
except Exception as e:
    logger.critical(f"❌ CSV save failed: {e}", exc_info=True)
    sys.exit(1)

logger.info(f"✅ {company} scraper completed successfully")
sys.exit(0)